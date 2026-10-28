import csv
import hashlib
import io
import zipfile
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import PureWindowsPath

from django.core.exceptions import ValidationError
from django.core.validators import DecimalValidator
from django.db import transaction
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from rest_framework import status
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.calificaciones.models import (
    CargaMasiva,
    DetalleCarga,
    Instrumento,
    LogAuditoria,
    OutboxCarga,
)
from apps.cuentas.permissions import CargaMasivaPermission

from .calculo import ErrorCalculo, calcular_factores, validar_suma_factores
from .cargas_integridad import _claves_existentes, _hash_payload

TAMANO_MAXIMO = 50 * 1024 * 1024
FILAS_MAXIMAS = 10_000
FILAS_MUESTRA = 20
ERRORES_MAXIMOS = 100
CAMPOS_EVENTO = (
    "mercado",
    "codigo_instrumento",
    "ejercicio",
    "secuencia_evento",
    "numero_dividendo",
    "fecha_pago",
    "descripcion_dividendo",
)


def _columnas_tipo(tipo):
    prefijo = "monto" if tipo == "MONTOS" else "factor"
    return (*CAMPOS_EVENTO, *(f"{prefijo}_{numero}" for numero in range(8, 38)))


def _nombre_seguro(nombre):
    base = PureWindowsPath(nombre.replace("/", "\\")).name
    return "".join(caracter for caracter in base if caracter.isprintable())[:200]


def _hash_archivo(archivo):
    archivo.seek(0)
    digest = hashlib.sha256()
    for bloque in archivo.chunks():
        digest.update(bloque)
    archivo.seek(0)
    return digest.hexdigest()


def _leer_csv(archivo):
    archivo.seek(0)
    texto = io.TextIOWrapper(archivo.file, encoding="utf-8-sig", newline="")
    try:
        filas = []
        lector = csv.reader(texto)
        for numero, fila in enumerate(lector):
            if numero > FILAS_MAXIMAS:
                raise ValueError(
                    f"El archivo supera el máximo de {FILAS_MAXIMAS} filas."
                )
            filas.append(fila)
        return filas
    finally:
        texto.detach()


def _validar_paquete_xlsx(archivo):
    archivo.seek(0)
    with zipfile.ZipFile(archivo) as paquete:
        miembros = paquete.infolist()
        if len(miembros) > 10_000:
            raise ValueError("El libro contiene demasiados componentes.")
        tamano_expandido = sum(miembro.file_size for miembro in miembros)
        if tamano_expandido > 200 * 1024 * 1024:
            raise ValueError("El contenido expandido del libro es demasiado grande.")
        for miembro in miembros:
            if miembro.file_size and not miembro.compress_size:
                raise ValueError("El libro contiene un componente comprimido inválido.")
            if (
                miembro.compress_size
                and miembro.file_size / miembro.compress_size > 200
            ):
                raise ValueError("El libro supera el límite de compresión permitido.")
    archivo.seek(0)


def _leer_xlsx(archivo):
    _validar_paquete_xlsx(archivo)
    libro = load_workbook(archivo, read_only=True, data_only=False)
    try:
        hoja = libro.active
        if hoja.max_row and hoja.max_row > FILAS_MAXIMAS + 1:
            raise ValueError(
                f"El archivo supera el máximo de {FILAS_MAXIMAS} filas."
            )
        if hoja.max_column and hoja.max_column > len(CAMPOS_EVENTO) + 30:
            raise ValueError("El libro contiene más columnas que la plantilla.")
        filas = [
            list(fila)
            for fila in hoja.iter_rows(
                min_row=1,
                max_row=min(hoja.max_row or 1, FILAS_MAXIMAS + 1),
                values_only=True,
            )
        ]
        return filas, hoja.title
    finally:
        libro.close()


def _validar_cabecera(fila, columnas_esperadas):
    if not fila:
        return None, {"detalle": "El archivo no contiene una fila de encabezados."}
    cabeceras = [
        str(valor).strip().lower() if valor is not None else ""
        for valor in fila
    ]
    if cabeceras and cabeceras[-1] == "":
        cabeceras.pop()
    duplicadas = sorted(
        {
            cabecera
            for cabecera in cabeceras
            if cabecera and cabeceras.count(cabecera) > 1
        }
    )
    faltantes = sorted(set(columnas_esperadas) - set(cabeceras))
    desconocidas = sorted(set(cabeceras) - set(columnas_esperadas))
    if duplicadas or faltantes or desconocidas:
        return None, {
            "detalle": "Los encabezados no coinciden con la plantilla seleccionada.",
            "duplicadas": duplicadas,
            "faltantes": faltantes,
            "desconocidas": desconocidas,
        }
    return {nombre: cabeceras.index(nombre) for nombre in cabeceras}, None


def _validar_entero(valor, minimo, maximo=None):
    if valor is None or str(valor).strip() == "":
        raise ValueError("El valor es obligatorio.")
    try:
        numero = Decimal(str(valor).strip())
    except InvalidOperation:
        raise ValueError("Debe ser un número entero.") from None
    if not numero.is_finite() or numero != numero.to_integral_value():
        raise ValueError("Debe ser un número entero.")
    resultado = int(numero)
    if resultado < minimo or (maximo is not None and resultado > maximo):
        raise ValueError("El valor está fuera del rango permitido.")
    return resultado


def _validar_decimal(valor, tipo):
    if valor is None or str(valor).strip() == "":
        return None
    try:
        numero = Decimal(str(valor).strip())
    except InvalidOperation:
        raise ValueError("Debe ser un número decimal válido.") from None
    if not numero.is_finite():
        raise ValueError("Debe ser un número decimal finito.")
    precision = DecimalValidator(18, 2) if tipo == "MONTOS" else DecimalValidator(9, 8)
    try:
        precision(numero)
    except ValidationError:
        detalle = (
            "18 dígitos y 2 decimales"
            if tipo == "MONTOS"
            else "9 dígitos y 8 decimales"
        )
        raise ValueError(
            f"Supera la precisión permitida ({detalle})."
        ) from None
    if numero < 0:
        raise ValueError("No se permiten valores negativos.")
    if tipo == "FACTORES" and numero > 1:
        raise ValueError("El factor debe estar entre 0 y 1.")
    return numero


def _validar_fila(fila, indices, columnas, tipo, instrumentos):
    errores = []
    if len(fila) > len(columnas) and any(
        valor not in (None, "") for valor in fila[len(columnas):]
    ):
        errores.append("La fila contiene más celdas que encabezados.")
    if len(fila) < len(columnas):
        fila = [*fila, *([None] * (len(columnas) - len(fila)))]
    valores = {
        nombre: fila[indice] if indice < len(fila) else None
        for nombre, indice in indices.items()
    }

    mercado = str(valores.get("mercado") or "").strip().upper()
    codigo = str(valores.get("codigo_instrumento") or "").strip().upper()
    instrumento_id = None
    ejercicio = None
    secuencia = None
    numero_dividendo = None
    fecha_pago = None

    if not mercado:
        errores.append("mercado: es obligatorio.")
    if not codigo:
        errores.append("codigo_instrumento: es obligatorio.")
    elif mercado:
        instrumento_id = instrumentos.get((mercado, codigo))
        if instrumento_id is None:
            errores.append(
                "El instrumento no existe o no está activo en ese mercado."
            )
    try:
        ejercicio = _validar_entero(valores.get("ejercicio"), 1900, 2200)
    except ValueError as error:
        errores.append(f"ejercicio: {error}")
    try:
        secuencia = _validar_entero(valores.get("secuencia_evento"), 1)
    except ValueError as error:
        errores.append(f"secuencia_evento: {error}")
    try:
        numero_dividendo = _validar_entero(valores.get("numero_dividendo"), 0)
    except ValueError as error:
        errores.append(f"numero_dividendo: {error}")

    fecha = valores.get("fecha_pago")
    try:
        if isinstance(fecha, datetime):
            fecha_pago = fecha.date()
        elif isinstance(fecha, date):
            fecha_pago = fecha
        else:
            fecha_pago = date.fromisoformat(str(fecha).strip())
    except (TypeError, ValueError):
        errores.append("fecha_pago: usa el formato AAAA-MM-DD.")

    descripcion = str(valores.get("descripcion_dividendo") or "").strip()
    if len(descripcion) > 200:
        errores.append("descripcion_dividendo: máximo 200 caracteres.")

    datos_normalizados = {
        "mercado": mercado,
        "codigo_instrumento": codigo,
        "ejercicio": ejercicio,
        "secuencia_evento": secuencia,
        "numero_dividendo": numero_dividendo,
        "fecha_pago": fecha_pago.isoformat() if fecha_pago else None,
        "descripcion_dividendo": descripcion,
    }
    for columna in columnas[len(CAMPOS_EVENTO):]:
        valor = valores.get(columna)
        if isinstance(valor, str) and valor.startswith("="):
            errores.append(f"{columna}: no se permiten fórmulas.")
            continue
        try:
            numero = _validar_decimal(valor, tipo)
            datos_normalizados[columna] = str(numero) if numero is not None else ""
        except ValueError as error:
            errores.append(f"{columna}: {error}")

    clave = None
    if (
        instrumento_id is not None
        and ejercicio is not None
        and secuencia is not None
        and fecha_pago is not None
    ):
        clave = (
            instrumento_id,
            ejercicio,
            fecha_pago,
            secuencia,
        )
    return datos_normalizados, errores, clave


def _contenido_fila_original(encabezados, fila):
    valores = [*fila, *([None] * max(0, len(encabezados) - len(fila)))]
    contenido = {}
    for encabezado, valor in zip(encabezados, valores):
        if isinstance(valor, (date, datetime)):
            valor = valor.isoformat()
        elif isinstance(valor, Decimal):
            valor = str(valor)
        elif valor is not None and not isinstance(valor, (str, int, float, bool)):
            valor = str(valor)
        contenido[str(encabezado)] = valor
    if len(fila) > len(encabezados):
        contenido["_celdas_adicionales"] = [
            (
                valor.isoformat()
                if isinstance(valor, (date, datetime))
                else str(valor) if valor is not None else None
            )
            for valor in fila[len(encabezados):]
        ]
    return contenido


class CargaMasivaPlantillaView(APIView):
    permission_classes = [IsAuthenticated, CargaMasivaPermission]

    def get(self, request):
        tipo = request.query_params.get("tipo", "").upper()
        formato = request.query_params.get("formato", "").upper()
        if tipo not in ("MONTOS", "FACTORES"):
            return Response(
                {"tipo": "Selecciona MONTOS o FACTORES."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        if formato not in ("CSV", "XLSX"):
            return Response(
                {"formato": "Selecciona CSV o XLSX."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        columnas = _columnas_tipo(tipo)
        nombre = f"plantilla_nuam_{tipo.lower()}.{formato.lower()}"
        if formato == "CSV":
            respuesta = HttpResponse(content_type="text/csv; charset=utf-8")
            respuesta.write("\ufeff")
            escritor = csv.writer(respuesta)
            escritor.writerow(columnas)
        else:
            libro = Workbook(write_only=True)
            hoja = libro.create_sheet("Carga")
            hoja.append(columnas)
            contenido = io.BytesIO()
            libro.save(contenido)
            respuesta = HttpResponse(
                contenido.getvalue(),
                content_type=(
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                ),
            )
        respuesta["Content-Disposition"] = f'attachment; filename="{nombre}"'
        return respuesta


class CargaMasivaPrevisualizarView(APIView):
    permission_classes = [IsAuthenticated, CargaMasivaPermission]

    def post(self, request):
        archivo = request.FILES.get("archivo")
        tipo = str(request.data.get("tipo", "")).upper()
        if tipo not in ("MONTOS", "FACTORES"):
            return Response(
                {"tipo": "Selecciona MONTOS o FACTORES."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        if archivo is None:
            return Response(
                {"archivo": "Selecciona un archivo CSV o XLSX."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        if not archivo.size or archivo.size > TAMANO_MAXIMO:
            return Response(
                {"archivo": "El archivo debe tener entre 1 byte y 50 MB."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        extension = PureWindowsPath(archivo.name.replace("/", "\\")).suffix.lower()
        if extension not in (".csv", ".xlsx"):
            return Response(
                {"archivo": "Solo se admiten archivos .csv y .xlsx."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        formato = extension[1:].upper()
        columnas = _columnas_tipo(tipo)
        try:
            if formato == "CSV":
                filas = _leer_csv(archivo)
                hoja = None
            else:
                filas, hoja = _leer_xlsx(archivo)
        except UnicodeDecodeError:
            return Response(
                {"archivo": "El CSV debe estar codificado en UTF-8."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        except ValueError as error:
            return Response(
                {"archivo": str(error)},
                status=status.HTTP_400_BAD_REQUEST,
            )
        except (
            csv.Error,
            InvalidFileException,
            KeyError,
            OSError,
            zipfile.BadZipFile,
        ):
            return Response(
                {"archivo": "No se pudo leer el archivo o no es un libro válido."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        if not filas:
            return Response(
                {"archivo": "El archivo está vacío."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        indices, error_cabecera = _validar_cabecera(filas[0], columnas)
        if error_cabecera:
            return Response(error_cabecera, status=status.HTTP_400_BAD_REQUEST)
        encabezados_detectados = [
            str(valor).strip() if valor is not None else ""
            for valor in filas[0]
        ]
        filas_datos = [
            fila
            for fila in filas[1:]
            if any(valor not in (None, "") for valor in fila)
        ]
        if not filas_datos:
            return Response(
                {"archivo": "Agrega al menos una fila con datos."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        if len(filas_datos) > FILAS_MAXIMAS:
            return Response(
                {"archivo": f"El máximo es de {FILAS_MAXIMAS} filas."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        instrumentos = {
            (item.mercado.codigo.upper(), item.codigo.upper()): item.pk
            for item in Instrumento.objects.filter(
                activo=True,
                mercado__activo=True,
            ).select_related("mercado")
        }
        resultados_filas = []
        claves_archivo = set()
        claves_validas = []
        for numero_fila, fila in enumerate(filas_datos, start=2):
            normalizada, errores_fila, clave = _validar_fila(
                fila,
                indices,
                columnas,
                tipo,
                instrumentos,
            )
            if tipo == "FACTORES" and not errores_fila:
                factores = {
                    numero: Decimal(
                        normalizada.get(f"factor_{numero}") or "0"
                    )
                    for numero in range(8, 38)
                }
                try:
                    validar_suma_factores(factores, desde=8, hasta=16)
                except ErrorCalculo as error:
                    errores_fila.append(str(error))
            if tipo == "MONTOS" and not errores_fila:
                montos = {
                    numero: Decimal(
                        normalizada.get(f"monto_{numero}") or "0"
                    )
                    for numero in range(8, 38)
                }
                try:
                    calcular_factores(montos, base_desde=8, base_hasta=19)
                except ErrorCalculo as error:
                    errores_fila.append(str(error))
            if clave and not errores_fila:
                if clave in claves_archivo:
                    errores_fila.append(
                        "La llave aparece más de una vez en el archivo."
                    )
                else:
                    claves_archivo.add(clave)
                    claves_validas.append(clave)
            resultados_filas.append(
                {
                    "fila": numero_fila,
                    "datos": normalizada,
                    "contenido_original": _contenido_fila_original(
                        encabezados_detectados,
                        fila,
                    ),
                    "errores": errores_fila,
                    "clave": clave,
                }
            )

        claves_en_base = _claves_existentes(
            claves_validas,
            request.nuam_profile,
        )
        errores = []
        muestra = []
        filas_validas = 0
        total_errores = 0
        for fila_resultado in resultados_filas:
            clave = fila_resultado.pop("clave")
            errores_fila = fila_resultado["errores"]
            if clave in claves_en_base and not errores_fila:
                errores_fila.append(
                    "Ya existe una calificación vigente con esta llave de negocio."
                )
            if errores_fila:
                total_errores += len(errores_fila)
                if len(errores) < ERRORES_MAXIMOS:
                    errores.extend(
                        {
                            "fila": fila_resultado["fila"],
                            "detalle": error,
                        }
                        for error in errores_fila[: ERRORES_MAXIMOS - len(errores)]
                    )
            else:
                filas_validas += 1
            if len(muestra) < FILAS_MUESTRA:
                muestra.append(
                    {
                        "fila": fila_resultado["fila"],
                        "datos": fila_resultado["datos"],
                        "errores": errores_fila,
                        "valida": not errores_fila,
                    }
                )

        filas_error = len(filas_datos) - filas_validas
        advertencias = [
            "La llave se compara con el estado actual y debe revalidarse "
            "al confirmar.",
            "Se valida la suma de factores 8–16 y, para montos, se comprueba "
            "el cálculo con base 8–19.",
            "Las celdas numéricas vacías permanecen vacías; no se convierten en cero.",
        ]
        advertencias.append(
            "Los nombres de las columnas 34–37 se tomaron del HDU; el tipo y "
            "la fórmula son metadatos provisionales pendientes de confirmación "
            "del docente."
        )
        carga = None
        operacion = "IMPORTAR" if filas_error == 0 else "VALIDACION_RECHAZOS"
        filas_outbox = [
            {
                "nro_fila": fila["fila"],
                "datos": fila["datos"],
                "contenido_original": fila["contenido_original"],
                "errores": fila["errores"],
            }
            for fila in resultados_filas
            if not fila["errores"] or operacion == "VALIDACION_RECHAZOS"
        ]
        payload = {
            "operacion": operacion,
            "tipo": tipo,
            "encabezados_detectados": encabezados_detectados,
            "hoja": hoja,
            "filas": filas_outbox,
        }
        sha256_payload = _hash_payload(payload)
        with transaction.atomic():
            carga = CargaMasiva.objects.create(
                corredor=request.nuam_profile.corredor,
                usuario=request.nuam_profile,
                tipo_carga=tipo,
                nombre_archivo=_nombre_seguro(archivo.name),
                formato=formato,
                hoja=hoja,
                encabezados_detectados=encabezados_detectados,
                tamano_bytes=archivo.size,
                sha256_archivo=_hash_archivo(archivo),
                total_filas=len(filas_datos),
                filas_ok=filas_validas,
                filas_error=filas_error,
                estado=(
                    CargaMasiva.Estado.PREVISUALIZADA
                    if filas_error == 0
                    else CargaMasiva.Estado.PENDIENTE_MONGO
                ),
                vencimiento_previsualizacion=(
                    timezone.now() + timedelta(minutes=30)
                    if filas_error == 0
                    else None
                ),
            )
            DetalleCarga.objects.bulk_create(
                [
                    DetalleCarga(
                        carga=carga,
                        nro_fila=fila["fila"],
                        estado_fila=(
                            DetalleCarga.EstadoFila.VALIDA
                            if not fila["errores"]
                            else DetalleCarga.EstadoFila.RECHAZADA
                        ),
                        mensaje_error="; ".join(fila["errores"])[:300],
                        contenido_original=fila["contenido_original"],
                        datos_normalizados=fila["datos"],
                    )
                    for fila in resultados_filas
                ],
                batch_size=500,
            )
            if filas_error:
                OutboxCarga.objects.create(
                    carga=carga,
                    payload=payload,
                    sha256_payload=sha256_payload,
                )
            LogAuditoria.objects.create(
                usuario=request.nuam_profile,
                corredor=request.nuam_profile.corredor,
                entidad="CargaMasiva",
                id_registro=carga.pk if carga else 0,
                accion=LogAuditoria.Accion.CARGA,
                valor_nuevo={
                    "tipo": tipo,
                    "formato": formato,
                    "total_filas": len(filas_datos),
                    "filas_ok": filas_validas,
                    "filas_error": filas_error,
                    "carga_id": carga.pk if carga else None,
                },
                metodo="PREVISUALIZACION",
            )
        if carga is not None:
            advertencias.append(
                "El archivo no se conserva. Los metadatos de la carga se "
                "registran en MongoDB; las filas válidas quedan como borrador "
                "temporal en PostgreSQL hasta confirmar o vencer."
                if filas_error == 0
                else (
                    "El archivo no se conserva. Las filas rechazadas y sus "
                    "errores se registran en MongoDB de forma asíncrona; las "
                    "filas válidas del mismo archivo no se aplican."
                )
            )
        return Response(
            {
                "archivo": _nombre_seguro(archivo.name),
                "tipo": tipo,
                "formato": formato,
                "carga_id": carga.pk if carga else None,
                "total_filas": len(filas_datos),
                "filas_validas": filas_validas,
                "filas_error": filas_error,
                "errores": errores,
                "errores_truncados": total_errores > len(errores),
                "muestra": muestra,
                "advertencias": advertencias,
                "confirmacion_habilitada": carga is not None and filas_error == 0,
            }
        )


def _respuesta_estado_carga(carga):
    outbox = getattr(carga, "outbox", None)
    if carga.estado == CargaMasiva.Estado.PENDIENTE_MONGO:
        mensaje = "Esperando confirmación del almacenamiento de detalles."
    elif carga.estado == CargaMasiva.Estado.CONFIRMADA:
        mensaje = "La carga se confirmó correctamente."
    elif carga.estado == CargaMasiva.Estado.FALLIDA:
        mensaje = (
            "La carga no se aplicó por un conflicto con registros existentes."
            if outbox and outbox.ultimo_error == "CONFLICTO_LLAVE"
            else "La carga no pudo completarse. Corrige el archivo y vuelve a cargar."
        )
    elif carga.estado == CargaMasiva.Estado.CANCELADA:
        mensaje = "La previsualización venció o fue cancelada."
    else:
        mensaje = "La carga está lista para confirmar."
    return {
        "id": carga.pk,
        "archivo": carga.nombre_archivo,
        "tipo": carga.tipo_carga,
        "estado": carga.estado,
        "total_filas": carga.total_filas,
        "filas_ok": carga.filas_ok,
        "filas_error": carga.filas_error,
        "fecha_hora": carga.fecha_hora,
        "vencimiento_previsualizacion": carga.vencimiento_previsualizacion,
        "intentos": outbox.intentos if outbox else 0,
        "mensaje": mensaje,
    }


def _cancelar_previsualizacion(carga, perfil, metodo):
    carga.estado = CargaMasiva.Estado.CANCELADA
    carga.save(update_fields=["estado"])
    carga.detalles.update(contenido_original={}, datos_normalizados={})
    LogAuditoria.objects.create(
        usuario=perfil,
        corredor=carga.corredor,
        entidad="CargaMasiva",
        id_registro=carga.pk,
        accion=LogAuditoria.Accion.CARGA,
        valor_nuevo={"estado": carga.estado},
        metodo=metodo,
    )


class CargaMasivaEstadoView(APIView):
    permission_classes = [IsAuthenticated, CargaMasivaPermission]

    def get(self, request, pk=None):
        cargas = CargaMasiva.objects.filter(usuario=request.nuam_profile)
        if pk is not None:
            carga = cargas.filter(pk=pk).first()
            if carga is None:
                return Response(
                    {"detalle": "No se encontró la carga."},
                    status=status.HTTP_404_NOT_FOUND,
                )
            return Response(_respuesta_estado_carga(carga))
        cargas = cargas.select_related("outbox")[:20]
        return Response([_respuesta_estado_carga(carga) for carga in cargas])

    def post(self, request, pk=None):
        if pk is None:
            return Response(
                {"detalle": "Debes indicar la carga que quieres confirmar."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        with transaction.atomic():
            carga = (
                CargaMasiva.objects.select_for_update()
                .filter(pk=pk, usuario=request.nuam_profile)
                .first()
            )
            if carga is None:
                return Response(
                    {"detalle": "No se encontró la carga."},
                    status=status.HTTP_404_NOT_FOUND,
                )
            if carga.estado == CargaMasiva.Estado.PENDIENTE_MONGO:
                return Response(
                    _respuesta_estado_carga(carga),
                    status=status.HTTP_202_ACCEPTED,
                )
            if carga.estado != CargaMasiva.Estado.PREVISUALIZADA:
                return Response(
                    {"detalle": "Esta carga ya no se puede confirmar."},
                    status=status.HTTP_409_CONFLICT,
                )
            if (
                carga.vencimiento_previsualizacion is None
                or carga.vencimiento_previsualizacion <= timezone.now()
            ):
                _cancelar_previsualizacion(
                    carga,
                    request.nuam_profile,
                    "VENCIMIENTO",
                )
                return Response(
                    {
                        "detalle": (
                            "La previsualización venció; vuelve a cargar el archivo."
                        )
                    },
                    status=status.HTTP_410_GONE,
                )

            detalles = list(carga.detalles.order_by("nro_fila"))
            if len(detalles) != carga.total_filas or not detalles:
                _cancelar_previsualizacion(
                    carga,
                    request.nuam_profile,
                    "BORRADOR_INCOMPLETO",
                )
                return Response(
                    {"detalle": "El borrador está incompleto y no puede confirmarse."},
                    status=status.HTTP_409_CONFLICT,
                )
            instrumentos = {
                (item.mercado.codigo.upper(), item.codigo.upper()): item.pk
                for item in Instrumento.objects.filter(
                    activo=True,
                    mercado__activo=True,
                ).select_related("mercado")
            }
            claves = []
            filas_payload = []
            for detalle in detalles:
                datos = detalle.datos_normalizados
                instrumento_id = instrumentos.get(
                    (
                        str(datos.get("mercado", "")).upper(),
                        str(datos.get("codigo_instrumento", "")).upper(),
                    )
                )
                if instrumento_id is None:
                    _cancelar_previsualizacion(
                        carga,
                        request.nuam_profile,
                        "INSTRUMENTO_NO_DISPONIBLE",
                    )
                    return Response(
                        {
                            "detalle": (
                                "Un instrumento dejó de estar disponible. El "
                                "borrador se canceló; vuelve a validar el archivo."
                            )
                        },
                        status=status.HTTP_409_CONFLICT,
                    )
                claves.append(
                    (
                        instrumento_id,
                        int(datos["ejercicio"]),
                        date.fromisoformat(datos["fecha_pago"]),
                        int(datos["secuencia_evento"]),
                    )
                )
                filas_payload.append(
                    {
                        "nro_fila": detalle.nro_fila,
                        "datos": datos,
                        "contenido_original": detalle.contenido_original,
                    }
                )
            if len(set(claves)) != len(claves):
                _cancelar_previsualizacion(
                    carga,
                    request.nuam_profile,
                    "LLAVES_DUPLICADAS",
                )
                return Response(
                    {"detalle": "El borrador contiene llaves de negocio duplicadas."},
                    status=status.HTTP_409_CONFLICT,
                )
            if _claves_existentes(claves, request.nuam_profile):
                _cancelar_previsualizacion(
                    carga,
                    request.nuam_profile,
                    "CONFLICTO_LLAVE",
                )
                return Response(
                    {
                        "detalle": (
                            "Ya existe una calificación con una de las llaves "
                            "del borrador. Se canceló el borrador; vuelve a "
                            "validar el archivo."
                        )
                    },
                    status=status.HTTP_409_CONFLICT,
                )

            payload = {
                "operacion": "IMPORTAR",
                "tipo": carga.tipo_carga,
                "encabezados_detectados": carga.encabezados_detectados,
                "hoja": carga.hoja,
                "filas": [
                    {**fila, "errores": []}
                    for fila in filas_payload
                ],
            }
            sha256_payload = _hash_payload(payload)
            OutboxCarga.objects.create(
                carga=carga,
                payload=payload,
                sha256_payload=sha256_payload,
            )
            carga.estado = CargaMasiva.Estado.PENDIENTE_MONGO
            carga.save(update_fields=["estado"])
            LogAuditoria.objects.create(
                usuario=request.nuam_profile,
                corredor=carga.corredor,
                entidad="CargaMasiva",
                id_registro=carga.pk,
                accion=LogAuditoria.Accion.CARGA,
                valor_nuevo={"estado": carga.estado, "sha256_payload": sha256_payload},
                metodo="CONFIRMAR",
            )
        carga.refresh_from_db()
        return Response(
            _respuesta_estado_carga(carga),
            status=status.HTTP_202_ACCEPTED,
        )
