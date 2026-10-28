import csv
import io
from datetime import date, timedelta

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework import status
from rest_framework.test import APIClient

from apps.calificaciones.cargas_api import CAMPOS_EVENTO, _columnas_tipo
from apps.calificaciones.models import (
    Calificacion,
    CargaMasiva,
    ColumnaTributaria,
    DetalleCarga,
    FactorCalificacion,
    Instrumento,
    LogAuditoria,
    Mercado,
    MontoCalificacion,
    OrigenInformacion,
    OutboxCarga,
)
from apps.cuentas.models import Corredor, Rol, Usuario
from apps.calificaciones.outbox import procesar_siguiente_carga
from pymongo.errors import ServerSelectionTimeoutError


class MongoCollectionFake:
    def __init__(self, database, nombre):
        self.database = database
        self.nombre = nombre
        self.documentos = {}
        self.fallar_escritura = False
        self.fallar_eliminacion = False

    @staticmethod
    def _clave(filtro):
        if "_id" in filtro:
            return filtro["_id"]
        if "fila" in filtro:
            return (filtro["id_carga"], filtro["fila"])
        return filtro["id_carga"]

    @staticmethod
    def _coincide(documento, filtro):
        for campo, valor in filtro.items():
            if isinstance(valor, dict) and "$in" in valor:
                if documento.get(campo) not in valor["$in"]:
                    return False
            elif documento.get(campo) != valor:
                return False
        return True

    def update_one(self, filtro, actualizacion, upsert=False):
        clave = self._clave(filtro)
        documento = self.documentos.get(clave)
        if documento is None and upsert:
            documento = dict(actualizacion.get("$setOnInsert", {}))
            for campo, valor in filtro.items():
                if not isinstance(valor, dict):
                    documento.setdefault(campo, valor)
            self.documentos[clave] = documento
        elif documento is not None:
            documento.update(actualizacion.get("$set", {}))

    def find_one(self, filtro, proyeccion):
        return next(
            (
                documento.copy()
                for documento in self.documentos.values()
                if self._coincide(documento, filtro)
            ),
            None,
        )

    def bulk_write(self, operaciones, ordered):
        if self.fallar_escritura:
            raise ServerSelectionTimeoutError("MongoDB no disponible")
        for operacion in operaciones:
            filtro = operacion._filter
            clave = self._clave(filtro)
            if clave not in self.documentos:
                documento = dict(operacion._doc["$setOnInsert"])
                for campo, valor in filtro.items():
                    if not isinstance(valor, dict):
                        documento.setdefault(campo, valor)
                self.documentos[clave] = documento

    def find(self, filtro, proyeccion):
        return [
            documento.copy()
            for documento in self.documentos.values()
            if self._coincide(documento, filtro)
        ]

    def delete_many(self, filtro):
        if self.fallar_eliminacion:
            raise ServerSelectionTimeoutError("MongoDB no disponible")
        for clave, documento in list(self.documentos.items()):
            if self._coincide(documento, filtro):
                del self.documentos[clave]


class MongoDatabaseFake:
    def __init__(self):
        self.colecciones = {}

    def __getitem__(self, nombre):
        if nombre not in self.colecciones:
            self.colecciones[nombre] = MongoCollectionFake(self, nombre)
        return self.colecciones[nombre]


class MongoClientFake:
    def __init__(self):
        self.database = MongoDatabaseFake()
        self.coleccion = self.database["calificaciones_detalle"]

    def __getitem__(self, nombre):
        return self.database


class CargasMasivasApiTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.rol_corredor = Rol.objects.create(
            codigo=Rol.Codigo.CORREDOR,
            nombre="Corredor",
        )
        cls.rol_operador = Rol.objects.create(
            codigo=Rol.Codigo.OPERADOR_BOLSA,
            nombre="Operador",
        )
        cls.rol_auditor = Rol.objects.create(
            codigo=Rol.Codigo.AUDITOR,
            nombre="Auditor",
        )
        cls.corredor = Corredor.objects.create(
            rut="76.555.555-5",
            nombre="Corredora de carga",
        )
        cls.otro_corredor = Corredor.objects.create(
            rut="76.666.666-6",
            nombre="Otra corredora",
        )
        cls.mercado = Mercado.objects.create(codigo="ACN", nombre="Acciones")
        cls.instrumento = Instrumento.objects.create(
            mercado=cls.mercado,
            codigo="COPEC",
            nombre="Copec",
        )
        cls.origen = OrigenInformacion.objects.create(
            codigo=OrigenInformacion.Codigo.SISTEMA,
            descripcion="Sistema",
        )
        OrigenInformacion.objects.create(
            codigo=OrigenInformacion.Codigo.ARCHIVO_CARGA,
            descripcion="Archivo de carga",
        )
        ColumnaTributaria.objects.bulk_create(
            [
                ColumnaTributaria(
                    nro_columna=numero,
                    nombre_homologado=f"Columna {numero}",
                )
                for numero in range(8, 38)
            ]
        )
        cls.operador = cls.crear_usuario(
            "operador-carga@example.test",
            cls.rol_operador,
        )
        cls.corredor_usuario = cls.crear_usuario(
            "corredor-carga@example.test",
            cls.rol_corredor,
            cls.corredor,
        )
        cls.otro_corredor_usuario = cls.crear_usuario(
            "otra-corredora-carga@example.test",
            cls.rol_corredor,
            cls.otro_corredor,
        )
        cls.auditor = cls.crear_usuario(
            "auditor-carga@example.test",
            cls.rol_auditor,
        )

    @classmethod
    def crear_usuario(cls, email, rol, corredor=None):
        cuenta = get_user_model().objects.create_user(
            username=email,
            email=email,
            password="Clave-de-prueba-2030!",
        )
        return Usuario.objects.create(
            cuenta=cuenta,
            email=email,
            rol=rol,
            corredor=corredor,
        )

    @staticmethod
    def cliente(perfil):
        client = APIClient()
        client.force_authenticate(user=perfil.cuenta)
        return client

    @staticmethod
    def fila(tipo="MONTOS"):
        columnas = _columnas_tipo(tipo)
        valores = {
            "mercado": "ACN",
            "codigo_instrumento": "COPEC",
            "ejercicio": "2025",
            "secuencia_evento": "10001",
            "numero_dividendo": "0",
            "fecha_pago": "2025-05-12",
            "descripcion_dividendo": "Dividendo de prueba",
        }
        prefijo = "monto" if tipo == "MONTOS" else "factor"
        valores[f"{prefijo}_8"] = "100.00" if tipo == "MONTOS" else "0.5"
        return [valores.get(columna, "") for columna in columnas]

    def archivo_csv(self, tipo="MONTOS", filas=None, columnas=None):
        contenido = io.StringIO(newline="")
        escritor = csv.writer(contenido)
        escritor.writerow(columnas or _columnas_tipo(tipo))
        for fila in filas or [self.fila(tipo)]:
            escritor.writerow(fila)
        return SimpleUploadedFile(
            f"datos.{tipo.lower()}.csv",
            contenido.getvalue().encode("utf-8"),
            content_type="text/csv",
        )

    def archivo_xlsx(self, tipo="FACTORES"):
        libro = Workbook()
        hoja = libro.active
        hoja.append(_columnas_tipo(tipo))
        hoja.append(self.fila(tipo))
        contenido = io.BytesIO()
        libro.save(contenido)
        return SimpleUploadedFile(
            "datos.xlsx",
            contenido.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )

    def test_plantillas_csv_y_xlsx_exponen_las_37_columnas(self):
        client = self.cliente(self.operador)
        csv_response = client.get(
            "/api/v1/cargas/plantilla/?tipo=MONTOS&formato=CSV"
        )
        xlsx_response = client.get(
            "/api/v1/cargas/plantilla/?tipo=FACTORES&formato=XLSX"
        )

        self.assertEqual(csv_response.status_code, status.HTTP_200_OK)
        self.assertIn("plantilla_nuam_montos.csv", csv_response["Content-Disposition"])
        self.assertEqual(xlsx_response.status_code, status.HTTP_200_OK)
        libro = load_workbook(io.BytesIO(xlsx_response.content), read_only=True)
        try:
            self.assertEqual(
                tuple(next(libro.active.iter_rows(values_only=True))),
                _columnas_tipo("FACTORES"),
            )
        finally:
            libro.close()
        self.assertEqual(len(_columnas_tipo("MONTOS")), 37)

    def test_previsualiza_csv_guarda_borrador_temporal_sin_confirmarlo(self):
        response = self.cliente(self.corredor_usuario).post(
            "/api/v1/cargas/previsualizar/",
            {"tipo": "MONTOS", "archivo": self.archivo_csv()},
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["total_filas"], 1)
        self.assertEqual(response.data["filas_validas"], 1)
        self.assertTrue(response.data["confirmacion_habilitada"])
        self.assertIsNotNone(response.data["carga_id"])
        self.assertTrue(
            any(
                "provisionales" in advertencia
                for advertencia in response.data["advertencias"]
            )
        )
        self.assertEqual(Calificacion.objects.count(), 0)
        self.assertEqual(CargaMasiva.objects.count(), 1)
        self.assertEqual(DetalleCarga.objects.count(), 1)
        self.assertEqual(
            CargaMasiva.objects.get().estado,
            CargaMasiva.Estado.PREVISUALIZADA,
        )
        self.assertTrue(
            LogAuditoria.objects.filter(
                usuario=self.corredor_usuario,
                accion=LogAuditoria.Accion.CARGA,
                metodo="PREVISUALIZACION",
            ).exists()
        )

    def test_error_de_datos_persiste_rechazo_pendiente_de_mongo(self):
        valores = self.fila("MONTOS")
        valores[7] = "0"
        response = self.cliente(self.operador).post(
            "/api/v1/cargas/previsualizar/",
            {
                "tipo": "MONTOS",
                "archivo": self.archivo_csv("MONTOS", [valores]),
            },
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["filas_validas"], 0)
        self.assertEqual(response.data["filas_error"], 1)
        self.assertFalse(response.data["confirmacion_habilitada"])
        carga = CargaMasiva.objects.get(pk=response.data["carga_id"])
        self.assertEqual(carga.estado, CargaMasiva.Estado.PENDIENTE_MONGO)
        self.assertTrue(OutboxCarga.objects.filter(carga=carga).exists())
        self.assertEqual(DetalleCarga.objects.filter(carga=carga).count(), 1)
        self.assertEqual(Calificacion.objects.count(), 0)

    def test_estado_y_lista_de_cargas_solo_exponen_al_propietario(self):
        carga_id = self._crear_borrador_confirmable()
        client = self.cliente(self.otro_corredor_usuario)

        listado = client.get("/api/v1/cargas/")
        detalle = client.get(f"/api/v1/cargas/{carga_id}/")
        confirmar = client.post(
            f"/api/v1/cargas/{carga_id}/confirmar/",
            {},
            format="json",
        )

        self.assertEqual(listado.status_code, status.HTTP_200_OK)
        self.assertEqual(listado.data, [])
        self.assertEqual(detalle.status_code, status.HTTP_404_NOT_FOUND)
        self.assertEqual(confirmar.status_code, status.HTTP_404_NOT_FOUND)
        self.assertEqual(
            CargaMasiva.objects.get(pk=carga_id).estado,
            CargaMasiva.Estado.PREVISUALIZADA,
        )

    def _crear_borrador_confirmable(self, perfil=None):
        respuesta = self.cliente(perfil or self.corredor_usuario).post(
            "/api/v1/cargas/previsualizar/",
            {"tipo": "MONTOS", "archivo": self.archivo_csv()},
            format="multipart",
        )
        self.assertEqual(respuesta.status_code, status.HTTP_200_OK)
        return respuesta.data["carga_id"]

    def test_confirma_y_procesa_outbox_idempotente_en_mongo_y_postgres(self):
        carga_id = self._crear_borrador_confirmable()
        respuesta = self.cliente(self.corredor_usuario).post(
            f"/api/v1/cargas/{carga_id}/confirmar/",
            {},
            format="json",
        )
        respuesta_repetida = self.cliente(self.corredor_usuario).post(
            f"/api/v1/cargas/{carga_id}/confirmar/",
            {},
            format="json",
        )
        carga = CargaMasiva.objects.get(pk=carga_id)
        outbox = OutboxCarga.objects.get(carga=carga)
        cliente_mongo = MongoClientFake()

        self.assertEqual(respuesta.status_code, status.HTTP_202_ACCEPTED)
        self.assertEqual(respuesta_repetida.status_code, status.HTTP_202_ACCEPTED)
        self.assertEqual(OutboxCarga.objects.filter(carga=carga).count(), 1)
        self.assertEqual(carga.estado, CargaMasiva.Estado.PENDIENTE_MONGO)
        self.assertEqual(outbox.estado, OutboxCarga.Estado.PENDIENTE)
        self.assertTrue(procesar_siguiente_carga(cliente_mongo))
        self.assertEqual(
            CargaMasiva.objects.get(pk=carga_id).estado,
            CargaMasiva.Estado.CONFIRMADA,
        )
        self.assertEqual(Calificacion.objects.count(), 1)
        self.assertEqual(MontoCalificacion.objects.count(), 1)
        self.assertEqual(FactorCalificacion.objects.count(), 30)
        self.assertEqual(len(cliente_mongo.coleccion.documentos), 1)
        self.assertFalse(procesar_siguiente_carga(cliente_mongo))
        self.assertEqual(Calificacion.objects.count(), 1)

    def test_reintenta_falla_de_mongo_sin_confirmar_el_lote(self):
        carga_id = self._crear_borrador_confirmable()
        self.cliente(self.corredor_usuario).post(
            f"/api/v1/cargas/{carga_id}/confirmar/",
            {},
            format="json",
        )
        cliente_mongo = MongoClientFake()
        cliente_mongo.coleccion.fallar_escritura = True

        self.assertTrue(procesar_siguiente_carga(cliente_mongo))
        carga = CargaMasiva.objects.get(pk=carga_id)
        outbox = OutboxCarga.objects.get(carga=carga)

        self.assertEqual(carga.estado, CargaMasiva.Estado.PENDIENTE_MONGO)
        self.assertEqual(outbox.estado, OutboxCarga.Estado.PENDIENTE)
        self.assertEqual(outbox.intentos, 1)
        self.assertEqual(outbox.ultimo_error, "ERROR_MONGO")
        self.assertEqual(Calificacion.objects.count(), 0)

    def test_registra_rechazos_en_mongo_sin_aplicar_filas_validas(self):
        columnas = _columnas_tipo("MONTOS")
        fila = self.fila()
        fila[columnas.index("monto_8")] = "-1"
        respuesta = self.cliente(self.operador).post(
            "/api/v1/cargas/previsualizar/",
            {
                "tipo": "MONTOS",
                "archivo": self.archivo_csv(filas=[fila]),
            },
            format="multipart",
        )
        carga_id = respuesta.data["carga_id"]
        cliente_mongo = MongoClientFake()

        self.assertEqual(respuesta.status_code, status.HTTP_200_OK)
        self.assertEqual(respuesta.data["filas_error"], 1)
        self.assertTrue(procesar_siguiente_carga(cliente_mongo))

        carga = CargaMasiva.objects.get(pk=carga_id)
        metadata = cliente_mongo.database["cargas_detalle"].find_one(
            {"id_carga": carga_id},
            {},
        )
        rechazadas = cliente_mongo.database[
            "cargas_filas_rechazadas"
        ].find({"id_carga": carga_id}, {})
        self.assertEqual(carga.estado, CargaMasiva.Estado.FALLIDA)
        self.assertEqual(carga.outbox.estado, OutboxCarga.Estado.PROCESADA)
        self.assertEqual(metadata["total_rechazadas"], 1)
        self.assertEqual(len(rechazadas), 1)
        self.assertEqual(rechazadas[0]["fila"], 2)
        self.assertTrue(rechazadas[0]["errores"])
        self.assertEqual(
            rechazadas[0]["contenido_original"]["monto_8"],
            "-1",
        )
        self.assertEqual(carga.detalles.get().contenido_original, {})
        self.assertEqual(carga.detalles.get().datos_normalizados, {})
        self.assertEqual(cliente_mongo.coleccion.documentos, {})
        self.assertEqual(Calificacion.objects.count(), 0)

    def test_conflicto_detectado_al_confirmar_cancela_el_borrador(self):
        carga_id = self._crear_borrador_confirmable()
        Calificacion.objects.create(
            corredor=self.corredor,
            instrumento=self.instrumento,
            origen=self.origen,
            usuario=self.corredor_usuario,
            ejercicio=2025,
            secuencia_evento=10001,
            fecha_pago=date(2025, 5, 12),
        )

        respuesta = self.cliente(self.corredor_usuario).post(
            f"/api/v1/cargas/{carga_id}/confirmar/",
            {},
            format="json",
        )
        carga = CargaMasiva.objects.get(pk=carga_id)

        self.assertEqual(respuesta.status_code, status.HTTP_409_CONFLICT)
        self.assertEqual(carga.estado, CargaMasiva.Estado.CANCELADA)
        self.assertFalse(OutboxCarga.objects.filter(carga=carga).exists())
        self.assertEqual(carga.detalles.get().contenido_original, {})
        self.assertEqual(carga.detalles.get().datos_normalizados, {})

    def test_worker_cancela_borradores_vencidos_y_limpia_datos_temporales(self):
        carga_id = self._crear_borrador_confirmable()
        CargaMasiva.objects.filter(pk=carga_id).update(
            vencimiento_previsualizacion=timezone.now() - timedelta(seconds=1)
        )

        self.assertFalse(procesar_siguiente_carga(MongoClientFake()))
        carga = CargaMasiva.objects.get(pk=carga_id)

        self.assertEqual(carga.estado, CargaMasiva.Estado.CANCELADA)
        self.assertEqual(carga.detalles.get().contenido_original, {})
        self.assertEqual(carga.detalles.get().datos_normalizados, {})

    def test_compensa_mongo_si_aparece_conflicto_de_llave_al_persistir(self):
        carga_id = self._crear_borrador_confirmable()
        self.cliente(self.corredor_usuario).post(
            f"/api/v1/cargas/{carga_id}/confirmar/",
            {},
            format="json",
        )
        Calificacion.objects.create(
            corredor=self.corredor,
            instrumento=self.instrumento,
            origen=self.origen,
            usuario=self.corredor_usuario,
            ejercicio=2025,
            secuencia_evento=10001,
            fecha_pago=date(2025, 5, 12),
        )
        cliente_mongo = MongoClientFake()

        self.assertTrue(procesar_siguiente_carga(cliente_mongo))
        self.assertEqual(
            OutboxCarga.objects.get(carga_id=carga_id).estado,
            OutboxCarga.Estado.COMPENSAR,
        )
        self.assertTrue(procesar_siguiente_carga(cliente_mongo))
        carga = CargaMasiva.objects.get(pk=carga_id)

        self.assertEqual(carga.estado, CargaMasiva.Estado.FALLIDA)
        self.assertEqual(carga.filas_error, 1)
        self.assertEqual(cliente_mongo.coleccion.documentos, {})
        self.assertEqual(Calificacion.objects.count(), 1)

    def test_reintenta_compensacion_sin_perder_el_motivo_original(self):
        carga_id = self._crear_borrador_confirmable()
        self.cliente(self.corredor_usuario).post(
            f"/api/v1/cargas/{carga_id}/confirmar/",
            {},
            format="json",
        )
        Calificacion.objects.create(
            corredor=self.corredor,
            instrumento=self.instrumento,
            origen=self.origen,
            usuario=self.corredor_usuario,
            ejercicio=2025,
            secuencia_evento=10001,
            fecha_pago=date(2025, 5, 12),
        )
        cliente_mongo = MongoClientFake()
        procesar_siguiente_carga(cliente_mongo)
        cliente_mongo.coleccion.fallar_eliminacion = True

        self.assertTrue(procesar_siguiente_carga(cliente_mongo))
        outbox = OutboxCarga.objects.get(carga_id=carga_id)
        self.assertEqual(outbox.estado, OutboxCarga.Estado.COMPENSAR)
        self.assertEqual(outbox.ultimo_error, "CONFLICTO_LLAVE")

        cliente_mongo.coleccion.fallar_eliminacion = False
        OutboxCarga.objects.filter(pk=outbox.pk).update(
            proximo_intento=timezone.now()
        )
        self.assertTrue(procesar_siguiente_carga(cliente_mongo))
        self.assertEqual(
            CargaMasiva.objects.get(pk=carga_id).estado,
            CargaMasiva.Estado.FALLIDA,
        )
        self.assertEqual(cliente_mongo.coleccion.documentos, {})

    def test_previsualiza_xlsx_y_rechaza_factores_fuera_de_rango(self):
        archivo = self.archivo_xlsx()
        libro = load_workbook(io.BytesIO(archivo.read()))
        libro.active["H2"] = "1.1"
        contenido = io.BytesIO()
        libro.save(contenido)
        archivo = SimpleUploadedFile("factores.xlsx", contenido.getvalue())

        response = self.cliente(self.operador).post(
            "/api/v1/cargas/previsualizar/",
            {"tipo": "FACTORES", "archivo": archivo},
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["filas_validas"], 0)
        self.assertEqual(response.data["filas_error"], 1)
        self.assertTrue(
            any("factor_8" in error["detalle"] for error in response.data["errores"])
        )

    def test_suma_de_factores_se_valida_en_columnas_8_a_16(self):
        valores = self.fila("FACTORES")
        valores[8] = "0.8"
        valores[9] = "0.5"
        archivo = self.archivo_csv("FACTORES", [valores])

        response = self.cliente(self.operador).post(
            "/api/v1/cargas/previsualizar/",
            {"tipo": "FACTORES", "archivo": archivo},
            format="multipart",
        )

        self.assertEqual(response.data["filas_error"], 1)
        self.assertIn("superar 1", response.data["errores"][0]["detalle"])

    def test_llave_existente_se_rechaza_solo_en_el_mismo_ambito(self):
        Calificacion.objects.create(
            corredor=self.corredor,
            instrumento=self.instrumento,
            origen=self.origen,
            usuario=self.corredor_usuario,
            ejercicio=2025,
            secuencia_evento=10001,
            fecha_pago=date(2025, 5, 12),
        )
        archivo_corredor = self.archivo_csv()
        archivo_operador = self.archivo_csv()

        respuesta_corredor = self.cliente(self.corredor_usuario).post(
            "/api/v1/cargas/previsualizar/",
            {"tipo": "MONTOS", "archivo": archivo_corredor},
            format="multipart",
        )
        respuesta_operador = self.cliente(self.operador).post(
            "/api/v1/cargas/previsualizar/",
            {"tipo": "MONTOS", "archivo": archivo_operador},
            format="multipart",
        )

        self.assertEqual(respuesta_corredor.data["filas_error"], 1)
        self.assertTrue(
            any(
                "llave de negocio" in error["detalle"]
                for error in respuesta_corredor.data["errores"]
            )
        )
        self.assertEqual(respuesta_operador.data["filas_validas"], 1)

    def test_rechaza_encabezados_incorrectos_archivos_no_admitidos_y_roles(self):
        columnas_incompletas = list(CAMPOS_EVENTO)
        cabecera_incorrecta = self.archivo_csv(
            columnas=columnas_incompletas,
            filas=[["ACN"] * len(columnas_incompletas)],
        )
        respuesta_cabecera = self.cliente(self.operador).post(
            "/api/v1/cargas/previsualizar/",
            {"tipo": "MONTOS", "archivo": cabecera_incorrecta},
            format="multipart",
        )
        archivo_xml = SimpleUploadedFile("datos.xml", b"<xml/>")
        respuesta_formato = self.cliente(self.operador).post(
            "/api/v1/cargas/previsualizar/",
            {"tipo": "MONTOS", "archivo": archivo_xml},
            format="multipart",
        )
        respuesta_rol = self.cliente(self.auditor).get(
            "/api/v1/cargas/plantilla/?tipo=MONTOS&formato=CSV"
        )

        self.assertEqual(
            respuesta_cabecera.status_code,
            status.HTTP_400_BAD_REQUEST,
        )
        self.assertEqual(respuesta_formato.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(respuesta_rol.status_code, status.HTTP_403_FORBIDDEN)
